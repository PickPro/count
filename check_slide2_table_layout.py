#!/usr/bin/env python3
"""Inspect slide 2's PowerPoint table geometry without exporting document text.

Requires: python-pptx
Put this script next to:
  Final Deck.pptx
  Final_Deck_3Q26_Compact.pptx
Run: python check_slide2_table_layout.py
Creates: slide2_layout_diagnosis.txt
"""
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
FILES = [HERE / 'Final Deck.pptx', HERE / 'Final_Deck_3Q26_Compact.pptx']
OUT = HERE / 'slide2_layout_diagnosis.txt'


def inch(x):
    return round(x / 914400, 4)


def report(path):
    lines = [f'FILE: {path.name}']
    if not path.is_file():
        return lines + ['MISSING FILE', '']
    pres = Presentation(path)
    if len(pres.slides) < 2:
        return lines + ['MISSING SLIDE 2', '']
    slide = pres.slides[1]
    lines.append(f'Slide size inches: {inch(pres.slide_width)} x {inch(pres.slide_height)}')
    candidates = []
    for ix, shape in enumerate(slide.shapes):
        if not shape.has_table:
            continue
        table = shape.table
        rows, cols = len(table.rows), len(table.columns)
        summary = (rows >= 6 and cols == 8 and
                   any('Level 1' in c.text for r in table.rows for c in r.cells))
        candidates.append((shape, summary))
        lines.append(
            f'Table #{ix}: rows={rows}, cols={cols}, '
            f'left={inch(shape.left)}, top={inch(shape.top)}, '
            f'width={inch(shape.width)}, height={inch(shape.height)}, '
            f'bottom={inch(shape.top + shape.height)}, summary={summary}'
        )
        if summary:
            lines.append('  Row heights inches: ' + ', '.join(str(inch(row.height)) for row in table.rows))
            lines.append('  Column widths inches: ' + ', '.join(str(inch(col.width)) for col in table.columns))
            for rix, row in enumerate(table.rows):
                fonts = []
                margin_values = set()
                auto_fit_modes = set()
                for cell in row.cells:
                    margin_values.add((inch(cell.margin_top), inch(cell.margin_bottom)))
                    auto_fit_modes.add(str(cell.text_frame.auto_size))
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.font.size:
                                fonts.append(round(run.font.size.pt, 1))
                lines.append(f'  Row {rix} explicit font pts: {sorted(set(fonts)) or "Inherited from theme"}; margins(top,bottom): {sorted(margin_values)}; auto_size: {sorted(auto_fit_modes)}')
    summaries = [shape for shape, is_summary in candidates if is_summary]
    if len(summaries) == 1:
        s = summaries[0]
        near = [shape for shape in slide.shapes
                if shape is not s and shape.top > s.top + 300000
                and shape.top < s.top + s.height + 1300000
                and shape.left < s.left + s.width and shape.left + shape.width > s.left]
        positions = sorted({inch(x.top) for x in near})[:20]
        lines.append('Nearby shape top positions (inches, may include overlapping design elements): ' + (', '.join(map(str, positions)) or 'None'))
    lines.append('')
    return lines


def main():
    lines = ['SLIDE 2 TABLE GEOMETRY CHECK (NO DOCUMENT TEXT EXPORTED)', '']
    for path in FILES:
        lines.extend(report(path))
    OUT.write_text('\n'.join(lines), encoding='utf-8')
    print('\n'.join(lines))
    print('Saved:', OUT)


if __name__ == '__main__':
    main()
