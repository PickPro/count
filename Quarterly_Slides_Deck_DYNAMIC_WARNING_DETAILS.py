import argparse
import copy
import logging
import re
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from io import BytesIO
from pptx.text.text import _Run
from pptx.oxml.ns import qn


# ==========================================================
# LOGGER
# ==========================================================

def setup_logger(logfile):
    logging.basicConfig(
        filename=logfile,
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s"
    )
    return logging.getLogger("QuarterlySlidesDeck")


# ==========================================================
# DUPLICATE SLIDE
# ==========================================================

def duplicate_slide(prs, source_slide):
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)

    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image_stream = BytesIO(shape.image.blob)
            new_slide.shapes.add_picture(
                image_stream,
                shape.left,
                shape.top,
                shape.width,
                shape.height
            )
        else:
            copied = copy.deepcopy(shape.element)
            new_slide.shapes._spTree.insert_element_before(
                copied,
                'p:extLst'
            )

    return new_slide


# ==========================================================
# REPLACE PLACEHOLDERS
# ==========================================================

def replace_text_in_slide(
    slide,
    title_text,
    metric_text,
    action_text
):
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame"):
            continue

        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "{{TITLE}}" in run.text:
                    run.text = str(title_text)

                elif "{{METRIC}}" in run.text:
                    run.text = str(metric_text)

                elif "{{ACTION}}" in run.text:
                    run.text = str(action_text)


# ==========================================================
# HIGHLIGHT WARNING ONLY ON GENERATED YELLOW DETAIL SLIDES
# ==========================================================
# Never color the warning-section separator/header slide.
# The number of metric records varies; no physical slide range is assumed.
WARNING_KEYWORD_COLOR = RGBColor(255, 192, 0)


def highlight_warning_in_metric_title(slide, title_text):
    """Color only 'Warning' in a generated yellow metric title.

    Requires the paragraph to contain the expected title text so that
    required-action paragraphs and other section labels remain unchanged.
    Preserves other text, run styling, font sizes, and shape coordinates.
    """
    expected = str(title_text or "").strip()
    if not expected or expected.lower() == "nan":
        return 0

    count = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            full_text = paragraph.text
            if expected not in full_text or not re.search(r"\bwarning\b", full_text, re.I):
                continue
            # The warning title normally occupies a single template run.
            # Split that run to keep exact formatting of surrounding words.
            for run in list(paragraph.runs):
                fragments = re.split(r"(\bWarning\b)", run.text, flags=re.IGNORECASE)
                if len(fragments) == 1:
                    continue

                original_xml = copy.deepcopy(run._r)
                run.text = fragments[0]
                anchor = run._r
                for fragment in fragments[1:]:
                    if not fragment:
                        continue
                    cloned = copy.deepcopy(original_xml)
                    t = cloned.find(qn("a:t"))
                    if t is None:
                        continue
                    t.text = fragment
                    anchor.addnext(cloned)
                    anchor = cloned
                    if fragment.lower() == "warning":
                        _Run(cloned, paragraph).font.color.rgb = WARNING_KEYWORD_COLOR
                        count += 1
    return count


# ==========================================================
# GENERATE DECK
# ==========================================================

def generate_quarterly_ppt(
    master_file,
    threshold_file,
    template_ppt,
    output_ppt,
    logger
):
    logger.info(f"Reading Master File: {master_file}")

    master_df = pd.read_excel(
        master_file,
        engine="openpyxl"
    )

    master_df.columns = (
        master_df.columns
        .astype(str)
        .str.strip()
    )

    logger.info(f"Reading Threshold File: {threshold_file}")

    threshold_df = pd.read_excel(
        threshold_file,
        sheet_name="Metrics and Thresholds",
        engine="openpyxl"
    )

    threshold_df.columns = (
        threshold_df.columns
        .astype(str)
        .str.strip()
    )

    master_df["File Name"] = (
        master_df["File Name"]
        .astype(str)
        .str.strip()
    )

    threshold_df["DownloadCrosstab"] = (
        threshold_df["DownloadCrosstab"]
        .astype(str)
        .str.strip()
    )

    threshold_df["Breach Slide Metric Name"] = (
        threshold_df["Breach Slide Metric Name"]
        .fillna("")  # preserve blanks so they are excluded below
        .astype(str)
        .str.strip()
    )

    # ==========================================================
    # RED METRICS
    # ==========================================================
    red_df = (
        master_df[
            master_df["Red"].notna()
        ][
            ["Slide No", "File Name"]
        ]
        .drop_duplicates()
        .sort_values(
            by=["Slide No", "File Name"]
        )
    )

    # ==========================================================
    # YELLOW METRICS
    # ==========================================================
    yellow_df = (
        master_df[
            master_df["Yellow"].notna()
        ][
            ["Slide No", "File Name"]
        ]
        .drop_duplicates()
        .sort_values(
            by=["Slide No", "File Name"]
        )
    )

    logger.info(
        f"Red Metrics Found: {len(red_df)}"
    )

    logger.info(
        f"Yellow Metrics Found: {len(yellow_df)}"
    )

    # ==========================================================
    # RED LOOKUP
    # ==========================================================
    red_output = red_df.merge(
        threshold_df[
            [
                "DownloadCrosstab",
                "Breach Slide Metric Name",
                "Limit Breach Title",
                "Limit Breach Required Action"
            ]
        ],
        left_on="File Name",
        right_on="DownloadCrosstab",
        how="left"
    )

    # ==========================================================
    # YELLOW LOOKUP
    # ==========================================================
    yellow_output = yellow_df.merge(
        threshold_df[
            [
                "DownloadCrosstab",
                "Breach Slide Metric Name",
                "Warning Level Title",
                "Warning Level Required Action"
            ]
        ],
        left_on="File Name",
        right_on="DownloadCrosstab",
        how="left"
    )

    # ==========================================================
    # REMOVE N/A / BLANK BREACH METRICS
    # ==========================================================
    red_before = len(red_output)
    yellow_before = len(yellow_output)

    red_output = red_output[
        (
            red_output["Breach Slide Metric Name"]
            .fillna("")
            .astype(str)
            .str.strip()
            .ne("")
        )
        &
        (
            red_output["Breach Slide Metric Name"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .ne("N/A")
        )
    ]

    yellow_output = yellow_output[
        (
            yellow_output["Breach Slide Metric Name"]
            .fillna("")
            .astype(str)
            .str.strip()
            .ne("")
        )
        &
        (
            yellow_output["Breach Slide Metric Name"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .ne("N/A")
        )
    ]

    logger.info(
        f"Red metrics excluded due to N/A/blank: "
        f"{red_before - len(red_output)}"
    )

    logger.info(
        f"Yellow metrics excluded due to N/A/blank: "
        f"{yellow_before - len(yellow_output)}"
    )

    # ==========================================================
    # OPEN TEMPLATE
    # ==========================================================
    logger.info(
        f"Opening Template PPT: {template_ppt}"
    )

    template_prs = Presentation(
        template_ppt
    )

    if len(template_prs.slides) < 4:
        raise Exception(
            "Template PPT must contain at least 4 slides."
        )

    limit_header_slide = template_prs.slides[0]
    limit_detail_slide = template_prs.slides[1]

    warning_header_slide = template_prs.slides[2]
    warning_detail_slide = template_prs.slides[3]

    # ==========================================================
    # OUTPUT DECK
    # ==========================================================
    output_prs = Presentation(
        template_ppt
    )

    while len(output_prs.slides) > 0:
        r_id = output_prs.slides._sldIdLst[0].rId

        output_prs.part.drop_rel(r_id)

        del output_prs.slides._sldIdLst[0]

    # ==========================================================
    # RED SECTION
    # ==========================================================
    if not red_output.empty:

        duplicate_slide(
            output_prs,
            limit_header_slide
        )

        for _, row in red_output.iterrows():

            slide = duplicate_slide(
                output_prs,
                limit_detail_slide
            )

            replace_text_in_slide(
                slide,
                row.get(
                    "Limit Breach Title",
                    ""
                ),
                row.get(
                    "Breach Slide Metric Name",
                    ""
                ),
                row.get(
                    "Limit Breach Required Action",
                    ""
                )
            )

    # ==========================================================
    # YELLOW SECTION
    # ==========================================================
    highlighted_words = 0
    yellow_title_slides = 0

    if not yellow_output.empty:

        duplicate_slide(
            output_prs,
            warning_header_slide
        )

        for _, row in yellow_output.iterrows():

            slide = duplicate_slide(
                output_prs,
                warning_detail_slide
            )

            warning_title = row.get("Warning Level Title", "")
            replace_text_in_slide(
                slide,
                warning_title,
                row.get("Breach Slide Metric Name", ""),
                row.get("Warning Level Required Action", "")
            )
            # Format during generation, and ONLY on this yellow detail slide.
            colored = highlight_warning_in_metric_title(slide, warning_title)
            if re.search(r"\bwarning\b", str(warning_title), re.I):
                if not colored:
                    raise ValueError(
                        "Could not color 'Warning' in generated yellow "
                        f"metric title: {warning_title!r}. "
                        "Check the PowerPoint template title placeholder."
                    )
                yellow_title_slides += 1
            highlighted_words += colored

    logger.info(
        f"Warning keyword formatted: {highlighted_words} occurrence(s) "
        f"in {yellow_title_slides} generated yellow metric slide(s). "
        "Warning section header and red slides remain unchanged."
    )

    # ==========================================================
    # SAVE
    # ==========================================================
    output_prs.save(output_ppt)

    logger.info(
        f"Created PPT: {output_ppt}"
    )

    logger.info(
        f"Red Slides: {len(red_output)}"
    )

    logger.info(
        f"Yellow Slides: {len(yellow_output)}"
    )

    logger.info(
        f"Total Slides: {len(output_prs.slides)}"
    )


# ==========================================================
# MAIN
# ==========================================================

def main():
    parser = argparse.ArgumentParser(
        description="Generate Quarterly Breach Deck"
    )

    parser.add_argument(
        "master_metrics_file"
    )

    parser.add_argument(
        "threshold_file"
    )

    parser.add_argument(
        "template_ppt"
    )

    parser.add_argument(
        "output_ppt"
    )

    parser.add_argument(
        "log_file"
    )

    parser.add_argument(
        "flag_file"
    )

    return parser.parse_args()


# ==========================================================
# ENTRY
# ==========================================================

if __name__ == "__main__":

    args = main()

    logger = setup_logger(
        args.log_file
    )

    try:

        logger.info(
            "Program Started"
        )

        generate_quarterly_ppt(
            args.master_metrics_file,
            args.threshold_file,
            args.template_ppt,
            args.output_ppt,
            logger
        )

        with open(
            args.flag_file,
            "w"
        ) as file:

            file.write("Success")

        logger.info(
            "Program Completed Successfully"
        )

    except Exception as ex:

        logger.exception(
            "Program Failed"
        )

        with open(
            args.flag_file,
            "w"
        ) as file:

            file.write(
                f"Error: {ex}"
            )

        raise
