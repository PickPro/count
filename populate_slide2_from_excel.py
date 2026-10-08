#!/usr/bin/env python3
r"""
Update the TOP-RIGHT metric summary on slide 2 using the generated Excel summary.

Source Excel (from the existing metric classification script):
  LatestFile_with_levels_and_summary.xlsx  ->  'Color Level Summary'

Target deck:
  Final Deck.pptx

Changes ONLY in the detected native PowerPoint summary table:
  * Blank three previous-quarter header cells (1Q26) and previous values.
  * Change three latest-quarter headers (2Q26 -> 3Q26).
  * Populate latest-quarter Green/Yellow/Red by Level 1/2/3, plus totals.
  * Update right-hand totals.
  * Keep the original deck and all other slides/tables unchanged.

The script deliberately stops rather than guessing when the summary isn't a
recognizable editable PowerPoint table. See the generated audit TXT file.

Install on a machine where approved:
    python -m pip install openpyxl python-pptx

Place this script next to Final Deck.pptx and the generated Excel, then run:
    python populate_slide2_from_excel.py

For A360 (Python Script > Open: Manual input, then Execute script):
    import runpy
    runpy.run_path(r'K:\CLE06\EBF_BOT_PROCESSING_TEST\RMCreditRiskCommitteeReporting\Scripts\populate_slide2_from_excel.py', run_name='__main__')

Optional flags:
    --excel PATH --pptx PATH --output PATH --slide 2 --inspect
r"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    raise SystemExit("Missing openpyxl. Ask for approval, then: python -m pip install openpyxl")

try:
    from pptx import Presentation
except ModuleNotFoundError:
    raise SystemExit("Missing python-pptx. Ask for approval, then: python -m pip install python-pptx")

QUARTER = "3Q26"
SUMMARY_SHEET = "Color Level Summary"
SOURCE_FILE = "LatestFile_with_levels_and_summary.xlsx"
PPTX_FILE = "Final Deck.pptx"
OUT_FILE = "Final_Deck_3Q26_Updated.pptx"
LEVELS = ("Level 1", "Level 2", "Level 3")
COLORS = ("Green", "Yellow", "Red")


def norm(x):
    return re.sub(r"\s+", " ", str(x or "")).strip().lower()


def clean_num(x, label):
    if x is None or isinstance(x, bool):
        raise ValueError(f"No numeric count for {label}; check {SUMMARY_SHEET}.")
    if isinstance(x, int):
        result = x
    elif isinstance(x, float) and x.is_integer():
        result = int(x)
    elif isinstance(x, str) and re.fullmatch(r"\d+", x.strip()):
        result = int(x.strip())
    else:
        raise ValueError(f"Invalid integer count for {label}: {x!r}")
    if result < 0:
        raise ValueError(f"Negative count for {label}: {result}")
    return result


def load_counts(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SUMMARY_SHEET not in wb.sheetnames:
            raise ValueError(f"Missing Excel sheet: {SUMMARY_SHEET}")
        ws = wb[SUMMARY_SHEET]
        vals = list(ws.iter_rows(min_row=1, max_row=100, max_col=15, values_only=True))
        header = None
        for ix, row in enumerate(vals):
            hdr = {norm(cell): col for col, cell in enumerate(row) if cell is not None}
            if set(("level", "green", "yellow", "red", "total")) <= set(hdr):
                header = (ix, hdr)
                break
        if header is None:
            raise ValueError("Summary headers Level/Green/Yellow/Red/Total not found")
        start, hdr = header
        seen = {}
        for row in vals[start + 1:]:
            label = str(row[hdr["level"]] or "").strip()
            if label not in (*LEVELS, "Grand Total"):
                continue
            if label in seen:
                raise ValueError(f"Duplicate summary row: {label}")
            r = {key: clean_num(row[hdr[key.lower()]], f"{label}/{key}")
                 for key in (*COLORS, "Total")}
            if r["Total"] != sum(r[c] for c in COLORS):
                raise ValueError(f"Totals do not reconcile for {label}: {r}")
            seen[label] = r
        absent = set((*LEVELS, "Grand Total")) - set(seen)
        if absent:
            raise ValueError(f"Missing required summary rows: {sorted(absent)}")
        for key in (*COLORS, "Total"):
            if sum(seen[lv][key] for lv in LEVELS) != seen["Grand Total"][key]:
                raise ValueError(f"Grand Total for {key} doesn't equal level totals")
        return seen
    finally:
        wb.close()


def replace_keep_format(text_frame, value):
    """Change displayed text without replacing the original shape/cell formatting."""
    new_text = str(value)
    if not text_frame.paragraphs:
        text_frame.text = new_text
        return
    first_para = text_frame.paragraphs[0]
    if first_para.runs:
        first_para.runs[0].text = new_text
        for run in first_para.runs[1:]:
            run.text = ""
    else:
        first_para.text = new_text
    for para in text_frame.paragraphs[1:]:
        for run in para.runs:
            run.text = ""


def row_kind(text):
    t = norm(text)
    m = re.fullmatch(r"level\s*([123])(?:\s*metric)?", t)
    if m:
        return f"Level {m.group(1)}"
    if t in ("total", "grand total"):
        return "Grand Total"
    return None


def extract_layout(slide, audit_path):
    lines = ["SLIDE 2 SHAPE / TABLE AUDIT", "If no editable summary table is found, send this report with the PPTX.", ""]
    for i, shape in enumerate(slide.shapes):
        location = f"x={shape.left/914400:.2f},y={shape.top/914400:.2f},w={shape.width/914400:.2f},h={shape.height/914400:.2f} in"
        lines.append(f"[{i}] name={shape.name!r}, type={str(shape.shape_type)!r}, {location}")
        if shape.has_table:
            tb = shape.table
            lines.append(f"  TABLE: rows={len(tb.rows)}, cols={len(tb.columns)}")
            for ri, row in enumerate(tb.rows):
                lines.append(f"    {ri}: " + " | ".join(repr(c.text[:70]) for c in row.cells))
        elif shape.has_text_frame:
            lines.append(f"  TEXT: {shape.text[:180]!r}")
    audit_path.write_text("\n".join(lines), encoding="utf-8")


def find_summary_table(slide):
    candidates = []
    for shape in slide.shapes:
        if not shape.has_table or len(shape.table.columns) != 8:
            continue
        t = shape.table
        rows = {}
        for idx, row in enumerate(t.rows):
            kind = row_kind(row.cells[0].text)
            if kind:
                rows[kind] = idx
        if all(lv in rows for lv in LEVELS) and "Grand Total" in rows:
            candidates.append((shape, rows))
    if len(candidates) != 1:
        raise ValueError(
            f"Expected exactly one 8-column editable summary table on slide 2; found {len(candidates)}. "
            "Check the slide layout audit and supply the PPTX for exact mapping."
        )
    return candidates[0]


def update_quarter_headers(slide, table_shape, previous="1Q26", current="2Q26"):
    changes = {previous.lower(): "", current.lower(): QUARTER}
    found_previous = found_current = 0
    t = table_shape.table

    # Header may be part of the table.
    for row in t.rows:
        for cell in row.cells:
            key = norm(cell.text)
            if key in changes:
                replace_keep_format(cell.text_frame, changes[key])
                if key == previous.lower():
                    found_previous += 1
                else:
                    found_current += 1

    # In some decks the quarter labels are separate text boxes above the table.
    # Only allow text boxes within the summary table's horizontal footprint,
    # and with a near top/bottom overlap (never edit the slide title).
    xl = table_shape.left
    xr = table_shape.left + table_shape.width
    yt = table_shape.top
    yb = table_shape.top + table_shape.height
    for shape in slide.shapes:
        if shape is table_shape or not shape.has_text_frame or shape.has_table:
            continue
        cx = shape.left + shape.width / 2
        cy = shape.top + shape.height / 2
        if not (xl <= cx <= xr and yt - table_shape.height * 0.35 <= cy <= yb):
            continue
        key = norm(shape.text)
        if key in changes:
            replace_keep_format(shape.text_frame, changes[key])
            if key == previous.lower():
                found_previous += 1
            else:
                found_current += 1
    if found_previous == 0 or found_current == 0:
        raise ValueError(
            f"Could not safely locate both previous/current quarter labels "
            f"(found previous={found_previous}, latest={found_current}). "
            "The upper summary may use separate grouped shapes. PPTX inspection is needed."
        )
    return found_previous, found_current


def update_pptx(pptx_file: Path, output_file: Path, counts, slide_number=2):
    pres = Presentation(str(pptx_file))
    if not (1 <= slide_number <= len(pres.slides)):
        raise ValueError(f"Slide {slide_number} not found: deck has {len(pres.slides)} slides")
    slide = pres.slides[slide_number - 1]
    audit_path = output_file.with_suffix(".audit.txt")
    extract_layout(slide, audit_path)
    shape, rows = find_summary_table(slide)
    table = shape.table

    # Note: summary has 8 logical columns (level, G-prev, G-new, Y-prev,
    # Y-new, R-prev, R-new, total). The source counts are NOT hardcoded.
    for label in (*LEVELS, "Grand Total"):
        excel_row = counts[label]
        ppt_row = rows[label]
        for col_index in (1, 3, 5):
            replace_keep_format(table.cell(ppt_row, col_index).text_frame, "")
        for col_index, color in ((2, "Green"), (4, "Yellow"), (6, "Red")):
            replace_keep_format(table.cell(ppt_row, col_index).text_frame, excel_row[color])
        replace_keep_format(table.cell(ppt_row, 7).text_frame, excel_row["Total"])

    previous_changes, current_changes = update_quarter_headers(slide, shape)
    pres.save(str(output_file))

    # Reload and verify that the exact values were written, including blanks.
    verify = Presentation(str(output_file))
    verified_shape, verified_rows = find_summary_table(verify.slides[slide_number - 1])
    verified = verified_shape.table
    for label in (*LEVELS, "Grand Total"):
        i = verified_rows[label]
        for k in (1, 3, 5):
            if verified.cell(i, k).text.strip():
                raise RuntimeError(f"Unexpected previous-quarter value remains at {label}/{k}")
        for j, color in ((2,"Green"),(4,"Yellow"),(6,"Red"),(7,"Total")):
            expected = counts[label][color]
            if verified.cell(i, j).text.strip() != str(expected):
                raise RuntimeError(f"Verification failure at {label}/{color}")
    return previous_changes, current_changes, audit_path


def main():
    default_dir = Path(__file__).resolve().parent
    p = argparse.ArgumentParser(description="Fill top-right slide 2 summary from generated Excel")
    p.add_argument("--excel", type=Path, default=default_dir / SOURCE_FILE)
    p.add_argument("--pptx", type=Path, default=default_dir / PPTX_FILE)
    p.add_argument("--output", type=Path, default=default_dir / OUT_FILE)
    p.add_argument("--slide", type=int, default=2)
    p.add_argument("--inspect", action="store_true", help="Only create audit of slide objects")
    args = p.parse_args()
    if not args.pptx.is_file():
        raise FileNotFoundError(f"PowerPoint not found: {args.pptx}")
    pres = Presentation(str(args.pptx))
    if not (1 <= args.slide <= len(pres.slides)):
        raise ValueError("Invalid slide number")
    if args.inspect:
        audit = args.output.with_suffix(".audit.txt")
        extract_layout(pres.slides[args.slide - 1], audit)
        print(f"Slide layout saved to: {audit}")
        return
    if not args.excel.is_file():
        raise FileNotFoundError(f"Generated Excel not found: {args.excel}")
    if args.pptx.resolve() == args.output.resolve():
        raise ValueError("Output file must be different to preserve the original presentation")
    counts = load_counts(args.excel)
    print(f"Reading {SUMMARY_SHEET} from {args.excel}")
    for label in (*LEVELS, "Grand Total"):
        x = counts[label]
        print(f"  {label:11}  Green={x['Green']:2}  Yellow={x['Yellow']:2}  Red={x['Red']:2}  Total={x['Total']:2}")
    prev, latest, audit = update_pptx(args.pptx, args.output, counts, slide_number=args.slide)
    print(f"Done. Previous-quarter headers cleared: {prev}, latest-quarter headers updated: {latest}")
    print(f"Updated PowerPoint: {args.output}")


if __name__ == "__main__":
    try:
        main()
    except Exception as ex:
        print(f"ERROR: {ex}", file=sys.stderr)
        raise
