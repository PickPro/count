#!/usr/bin/env python3
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    print("\nERROR: openpyxl is not installed.")
    print("Run this once:")
    print("python -m pip install openpyxl")
    sys.exit(1)

# ============================================================
# HARDCODED FILE PATHS
# ============================================================
FIRST_FILE = Path(r"C:\VsCode\Test\COUNT\LatestFile.xlsx")
SECOND_FILE = Path(r"C:\VsCode\Test\COUNT\metrics.xlsx")
OUTPUT_FILE = Path(r"C:\VsCode\Test\COUNT\LatestFile_with_levels_and_summary.xlsx")


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_name(value):
    text = clean(value)
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"\s*-\s*", " - ", text)
    text = re.sub(r"\s+", " ", text).strip()

    aliases = {
        "Constr Mat % of CET1 Graph": "Const Mat % of CET1 Graph",
        "Constr Mat Crit % Graph": "Const Mat Crit % Graph",
    }
    return aliases.get(text, text)


def normalize_level(value):
    text = clean(value)
    match = re.search(r"\bLevel\s*([123])\b", text, flags=re.IGNORECASE)
    return f"Level {match.group(1)}" if match else None


def has_value(value):
    return value is not None and clean(value) != ""


def find_sheet_and_headers(workbook, required_headers):
    required = {h.lower(): h for h in required_headers}

    for ws in workbook.worksheets:
        for row_num in range(1, min(ws.max_row, 30) + 1):
            found = {}

            for col_num in range(1, ws.max_column + 1):
                value = clean(ws.cell(row=row_num, column=col_num).value)
                key = value.lower()

                if key in required:
                    found[required[key]] = col_num

            if all(header in found for header in required_headers):
                return ws, row_num, found

    raise ValueError(
        "Required columns not found: " + ", ".join(required_headers)
    )


def main():
    print()
    print("Using files:")
    print(f"First file : {FIRST_FILE}")
    print(f"Second file: {SECOND_FILE}")
    print()

    if not FIRST_FILE.exists():
        print(f"ERROR: First file not found:\n{FIRST_FILE}")
        sys.exit(1)

    if not SECOND_FILE.exists():
        print(f"ERROR: Second file not found:\n{SECOND_FILE}")
        sys.exit(1)

    wb1 = load_workbook(FIRST_FILE)
    wb2 = load_workbook(SECOND_FILE, data_only=True)

    ws1, header1, cols1 = find_sheet_and_headers(
        wb1, ["File Name", "Green", "Yellow", "Red"]
    )
    ws2, header2, cols2 = find_sheet_and_headers(
        wb2, ["DownloadCrosstab", "GRC Metric Type"]
    )

    level_lookup = {}

    for row in range(header2 + 1, ws2.max_row + 1):
        metric = normalize_name(
            ws2.cell(row=row, column=cols2["DownloadCrosstab"]).value
        )
        level = clean(
            ws2.cell(row=row, column=cols2["GRC Metric Type"]).value
        )

        if metric:
            level_lookup[metric] = level

    level_col = None

    for col in range(1, ws1.max_column + 1):
        if clean(ws1.cell(row=header1, column=col).value).lower() == "level":
            level_col = col
            break

    if level_col is None:
        level_col = ws1.max_column + 1
        ws1.cell(row=header1, column=level_col).value = "Level"

    matched = 0
    unmatched = []

    for row in range(header1 + 1, ws1.max_row + 1):
        file_name = ws1.cell(row=row, column=cols1["File Name"]).value

        if not clean(file_name):
            continue

        key = normalize_name(file_name)
        level = level_lookup.get(key)

        if level is None:
            unmatched.append((row, clean(file_name)))
        else:
            ws1.cell(row=row, column=level_col).value = level
            matched += 1

    if unmatched:
        print("UNMATCHED METRICS")
        print("-" * 60)

        for row, name in unmatched:
            print(f"Row {row}: {name}")

        print()
        print(f"Total unmatched: {len(unmatched)}")
        print("Output was NOT saved because some metrics were unmatched.")
        sys.exit(1)

    counts = {
        "Level 1": {"Green": 0, "Yellow": 0, "Red": 0},
        "Level 2": {"Green": 0, "Yellow": 0, "Red": 0},
        "Level 3": {"Green": 0, "Yellow": 0, "Red": 0},
    }

    errors = []

    for row in range(header1 + 1, ws1.max_row + 1):
        file_name = clean(
            ws1.cell(row=row, column=cols1["File Name"]).value
        )

        if not file_name:
            continue

        level = normalize_level(
            ws1.cell(row=row, column=level_col).value
        )

        if level not in counts:
            errors.append(f"Row {row}: invalid level")
            continue

        colors = []

        if has_value(ws1.cell(row=row, column=cols1["Green"]).value):
            colors.append("Green")
        if has_value(ws1.cell(row=row, column=cols1["Yellow"]).value):
            colors.append("Yellow")
        if has_value(ws1.cell(row=row, column=cols1["Red"]).value):
            colors.append("Red")

        if len(colors) == 1:
            counts[level][colors[0]] += 1
        elif len(colors) == 0:
            errors.append(f"Row {row}: no Green/Yellow/Red value")
        else:
            errors.append(f"Row {row}: multiple colors found {colors}")

    if errors:
        print("DATA ERRORS")
        print("-" * 60)
        for error in errors:
            print(error)
        print()
        print("Output was NOT saved because data errors were found.")
        sys.exit(1)

    green_total = sum(counts[level]["Green"] for level in counts)
    yellow_total = sum(counts[level]["Yellow"] for level in counts)
    red_total = sum(counts[level]["Red"] for level in counts)
    grand_total = green_total + yellow_total + red_total

    summary_name = "Color Level Summary"

    if summary_name in wb1.sheetnames:
        del wb1[summary_name]

    summary = wb1.create_sheet(summary_name)
    summary.append(["Level", "Green", "Yellow", "Red", "Total"])

    for level in ["Level 1", "Level 2", "Level 3"]:
        green = counts[level]["Green"]
        yellow = counts[level]["Yellow"]
        red = counts[level]["Red"]
        summary.append([level, green, yellow, red, green + yellow + red])

    summary.append(
        ["Grand Total", green_total, yellow_total, red_total, grand_total]
    )

    wb1.save(OUTPUT_FILE)

    print()
    print("=" * 65)
    print("METRIC COLOR COUNT BY LEVEL")
    print("=" * 65)
    print(
        f"{'Level':<15}"
        f"{'Green':>10}"
        f"{'Yellow':>10}"
        f"{'Red':>10}"
        f"{'Total':>10}"
    )
    print("-" * 65)

    for level in ["Level 1", "Level 2", "Level 3"]:
        green = counts[level]["Green"]
        yellow = counts[level]["Yellow"]
        red = counts[level]["Red"]
        total = green + yellow + red

        print(
            f"{level:<15}"
            f"{green:>10}"
            f"{yellow:>10}"
            f"{red:>10}"
            f"{total:>10}"
        )

    print("-" * 65)
    print(
        f"{'Grand Total':<15}"
        f"{green_total:>10}"
        f"{yellow_total:>10}"
        f"{red_total:>10}"
        f"{grand_total:>10}"
    )
    print("=" * 65)
    print(f"Matched metrics : {matched}")
    print(f"Output file     : {OUTPUT_FILE}")
    print("=" * 65)
    print()


if __name__ == "__main__":
    main()
